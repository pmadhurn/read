"""Builds small EPUB / PDF / TXT / DOCX / HTML / MD fixtures (English, Hindi, Gujarati)."""
import sys, zipfile
from pathlib import Path

out = Path(sys.argv[1]); out.mkdir(parents=True, exist_ok=True)
LOREM = ("The quick brown fox jumps over the lazy dog, again and again; it never tires. "
         "Reading one word at a time feels strange at first! Then it becomes natural? Yes. ") * 40
HINDI = "यह एक परीक्षण पुस्तक है। क्षत्रिय और विद्यार्थी स्त्री के साथ पढ़ते हैं। ज्ञान ही शक्ति है। " * 30
GUJ = "આ એક પરીક્ષણ પુસ્તક છે. વિદ્યાર્થીઓ દરરોજ વાંચે છે. જ્ઞાન એ શક્તિ છે. " * 30

def chapter(n, body):
    return f"<html xmlns='http://www.w3.org/1999/xhtml'><head><title>c{n}</title></head><body><h1>Chapter {n}</h1>" + \
           "".join(f"<p>{body}</p>" for _ in range(3)) + "</body></html>"

with zipfile.ZipFile(out / "test.epub", "w") as z:
    z.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
    z.writestr("META-INF/container.xml", "<?xml version='1.0'?><container version='1.0' xmlns='urn:oasis:names:tc:opendocument:xmlns:container'><rootfiles><rootfile full-path='OEBPS/content.opf' media-type='application/oebps-package+xml'/></rootfiles></container>")
    items = "".join(f"<item id='c{i}' href='c{i}.xhtml' media-type='application/xhtml+xml'/>" for i in (1, 2, 3))
    refs = "".join(f"<itemref idref='c{i}'/>" for i in (1, 2, 3))
    z.writestr("OEBPS/content.opf", f"<?xml version='1.0'?><package xmlns='http://www.idpf.org/2007/opf' version='3.0'><metadata xmlns:dc='http://purl.org/dc/elements/1.1/'><dc:title>Test Book of Foxes</dc:title><dc:creator>Asha Author</dc:creator></metadata><manifest>{items}<item id='nav' href='nav.xhtml' media-type='application/xhtml+xml' properties='nav'/></manifest><spine>{refs}</spine></package>")
    z.writestr("OEBPS/nav.xhtml", "<html xmlns='http://www.w3.org/1999/xhtml' xmlns:epub='http://www.idpf.org/2007/ops'><body><nav epub:type='toc'><ol>" + "".join(f"<li><a href='c{i}.xhtml'>Chapter {i}: The {w}</a></li>" for i, w in ((1, 'Fox'), (2, 'Dog'), (3, 'End'))) + "</ol></nav></body></html>")
    z.writestr("OEBPS/c1.xhtml", chapter(1, LOREM))
    z.writestr("OEBPS/c2.xhtml", chapter(2, HINDI))
    z.writestr("OEBPS/c3.xhtml", chapter(3, GUJ))

(out / "test.txt").write_text("Chapter 1\n\n" + LOREM + "\n\nChapter 2\n\n" + LOREM, encoding="utf-8")
(out / "test.md").write_text("# Markdown Book\n\n## Part one\n\n" + LOREM + "\n\n## Part two\n\n" + LOREM, encoding="utf-8")
(out / "test.html").write_text("<html><head><title>Html Book</title></head><body><h1>Chapter 1</h1><p>" + LOREM + "</p><h1>Chapter 2</h1><p>" + LOREM + "</p></body></html>", encoding="utf-8")

import fitz
doc = fitz.open()
for n in range(1, 9):
    page = doc.new_page()
    page.insert_text((72, 40), "My Running Header", fontsize=9)
    page.insert_text((290, 800), str(n), fontsize=9)
    if n in (1, 5):
        page.insert_text((72, 90), f"Chapter {1 if n == 1 else 2}", fontsize=18)
    text = ("This paragraph has a hyphen-\nated word and continues across lines without stopping until the "
            "very end of the page where it is cut off mid sentence and then picks up on the following page ") * 6
    page.insert_textbox(fitz.Rect(72, 110, 520, 780), text, fontsize=11)
doc.set_metadata({"title": "Pdf Test Book", "author": "Rohan Writer"})
doc.set_toc([[1, "Chapter One", 1], [1, "Chapter Two", 5]])
doc.save(out / "test.pdf")

scan = fitz.open(); p = scan.new_page()
pix = fitz.open(out / "test.pdf")[0].get_pixmap(dpi=150)
p.insert_image(p.rect, pixmap=pix); scan.save(out / "scanned.pdf")

locked = fitz.open(out / "test.pdf")
locked.save(out / "locked.pdf", encryption=fitz.PDF_ENCRYPT_AES_256, user_pw="x", owner_pw="y")

import docx
d = docx.Document(); d.core_properties.title = "Docx Book"; d.core_properties.author = "Word Smith"
for i in (1, 2):
    d.add_heading(f"Chapter {i}", 1)
    for _ in range(3): d.add_paragraph(LOREM)
d.save(out / "test.docx")
(out / "corrupt.epub").write_bytes(b"not a zip at all")
print("fixtures ok")
